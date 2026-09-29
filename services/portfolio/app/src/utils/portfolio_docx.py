"""
Сборка портфолио в формате Word (.docx) на python-docx.

Оформление повторяет сайт в «бумажном» варианте: чернильный индиго для текста
и заголовков, янтарь для акцентов, светлая янтарная заливка для цифр.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from src.constants.portfolio_export import CONTACTS, PROFILE

AVATAR_PATH = Path(__file__).resolve().parent.parent / "templates" / "export" / "avatar.jpg"

FONT = "Calibri"
INK = RGBColor(0x1B, 0x1E, 0x3A)
TEXT = RGBColor(0x2E, 0x31, 0x4A)
MUTED = RGBColor(0x6B, 0x6E, 0x85)
AMBER = RGBColor(0xC2, 0x5E, 0x06)
CORAL = RGBColor(0xD9, 0x48, 0x2B)
AMBER_HEX = "D97706"
LINE_HEX = "E5E2DA"
TINT_HEX = "FFF5E6"

CONTENT_WIDTH = Cm(17.4)  # A4 минус поля 1.8 см


@dataclass
class ExportExperience:
    start_date: date
    end_date: date | None
    is_current: bool
    position: str
    company: str
    responsibilities: str | None = None
    description: str | None = None


@dataclass
class ExportProject:
    title: str
    short_description: str | None = None
    description: str | None = None
    stack: str | None = None


@dataclass
class ExportAchievement:
    title: str
    category: str
    scope: str
    year: int | None = None
    organization: str | None = None
    result: str | None = None
    description: str | None = None


@dataclass
class ExportMetric:
    value: str
    label: str


@dataclass
class PortfolioExportData:
    lang: str
    experiences: list[ExportExperience] = field(default_factory=list)
    projects: list[ExportProject] = field(default_factory=list)
    achievements: list[ExportAchievement] = field(default_factory=list)
    metrics: list[ExportMetric] = field(default_factory=list)
    mentoring_lead: str | None = None


# ---------------------------------------------------------------- XML-помощники
# Порядок дочерних элементов в OOXML строгий: вставляем перед теми, что идут после по схеме

_PBDR_SUCCESSORS = (
    "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
    "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid",
    "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
    "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle",
    "w:rPr", "w:sectPr", "w:pPrChange",
)
_TBL_BORDERS_SUCCESSORS = ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption",
                           "w:tblDescription", "w:tblPrChange")
_TC_SHD_SUCCESSORS = ("w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark",
                      "w:headers", "w:cellIns", "w:cellDel", "w:cellMerge", "w:tcPrChange")
_TC_MAR_SUCCESSORS = _TC_SHD_SUCCESSORS[2:]
_RPR_SPACING_SUCCESSORS = ("w:w", "w:kern", "w:position", "w:sz", "w:szCs", "w:highlight", "w:u",
                           "w:effect", "w:bdr", "w:shd", "w:fitText", "w:vertAlign", "w:rtl", "w:cs",
                           "w:em", "w:lang", "w:eastAsianLayout", "w:specVanish", "w:oMath")

def _set_cell_shading(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.insert_element_before(shd, *_TC_SHD_SUCCESSORS)


def _set_cell_margins(cell, top=60, bottom=60, left=100, right=100) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for side, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    tc_pr.insert_element_before(margins, *_TC_MAR_SUCCESSORS)


def _remove_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)
    tbl_pr.insert_element_before(borders, *_TBL_BORDERS_SUCCESSORS)


def _set_col_widths(table, widths) -> None:
    table.autofit = False
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width
    for column, width in zip(table.columns, widths):
        column.width = width


def _paragraph_border(paragraph, edge: str, color: str, size: int = 6, space: int = 4) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.insert_element_before(borders, *_PBDR_SUCCESSORS)
    node = OxmlElement(f"w:{edge}")
    node.set(qn("w:val"), "single")
    node.set(qn("w:sz"), str(size))
    node.set(qn("w:space"), str(space))
    node.set(qn("w:color"), color)
    borders.append(node)


def _row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:cantSplit")
    tr_pr.append(node)


def _add_hyperlink(paragraph, text: str, url: str, color: RGBColor, size: float) -> None:
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                          is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), FONT)
    r_pr.append(fonts)
    col = OxmlElement("w:color")
    col.set(qn("w:val"), str(color))
    r_pr.append(col)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    r_pr.append(sz)
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def _add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    _style_run(run, size=8, color=MUTED)
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            node = OxmlElement("w:fldChar")
            node.set(qn("w:fldCharType"), kind)
        else:
            node = OxmlElement("w:instrText")
            node.set(qn("xml:space"), "preserve")
            node.text = text
        run._r.append(node)


def _style_run(run, *, size=None, bold=None, italic=None, color=None, spacing=None):
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    if spacing is not None:
        r_pr = run._r.get_or_add_rPr()
        node = OxmlElement("w:spacing")
        node.set(qn("w:val"), str(spacing))
        r_pr.insert_element_before(node, *_RPR_SPACING_SUCCESSORS)
    return run


def _para(container, text: str = "", *, size=10, bold=False, italic=False, color=TEXT,
          before=0, after=4, line=1.15, align=None, keep_next=False):
    paragraph = container.add_paragraph()
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep_next
    if align is not None:
        paragraph.alignment = align
    if text:
        _style_run(paragraph.add_run(text), size=size, bold=bold, italic=italic, color=color)
    return paragraph


def _clear_cell(cell):
    """В новой ячейке уже есть пустой абзац — используем его, а не плодим отступы."""
    return cell.paragraphs[0]


# ---------------------------------------------------------------- разделы

class _Builder:
    def __init__(self, data: PortfolioExportData):
        self.data = data
        self.t = PROFILE.get(data.lang, PROFILE["ru"])
        self.doc = Document()
        self.section_no = 0
        self._setup()

    def _setup(self):
        section = self.doc.sections[0]
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.top_margin = section.bottom_margin = Cm(1.6)
        section.left_margin = section.right_margin = Cm(1.8)
        section.footer_distance = Cm(0.8)

        normal = self.doc.styles["Normal"]
        normal.font.name = FONT
        normal.font.size = Pt(10)
        normal.font.color.rgb = TEXT
        r_fonts = normal.element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            r_fonts.set(qn(attr), FONT)

        zoom = self.doc.settings.element.find(qn("w:zoom"))
        if zoom is not None and zoom.get(qn("w:percent")) is None:
            zoom.set(qn("w:percent"), "100")

        core = self.doc.core_properties
        core.title = f"{self.t['name']} — {self.t['footer']}"
        core.author = self.t["name"]

        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _style_run(footer.add_run(f"{self.t['name']} · {self.t['footer']} · "), size=8, color=MUTED)
        _add_page_field(footer)

    # --- шапка
    def header(self):
        table = self.doc.add_table(rows=1, cols=2)
        _remove_table_borders(table)
        _set_col_widths(table, [Cm(3.6), CONTENT_WIDTH - Cm(3.6)])
        photo_cell, info_cell = table.rows[0].cells
        for cell in (photo_cell, info_cell):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell, 0, 0, 0, 120)

        if AVATAR_PATH.exists():
            _clear_cell(photo_cell).add_run().add_picture(str(AVATAR_PATH), width=Cm(3.2))

        name = _clear_cell(info_cell)
        name.paragraph_format.space_after = Pt(2)
        _style_run(name.add_run(self.t["name"]), size=26, bold=True, color=INK, spacing=-10)
        _para(info_cell, self.t["role"], size=11.5, color=AMBER, after=8)

        contacts = _para(info_cell, after=0)
        for index, (label, url) in enumerate(CONTACTS):
            if index:
                _style_run(contacts.add_run("   ·   "), size=9, color=MUTED)
            _add_hyperlink(contacts, label, url, INK, 9)

        rule = _para(self.doc, after=6, before=6)
        _paragraph_border(rule, "bottom", AMBER_HEX, size=12, space=1)

    # --- заголовок раздела: «01  ОПЫТ РАБОТЫ» с тонкой линией
    def heading(self, title: str):
        self.section_no += 1
        paragraph = _para(self.doc, before=16, after=8, keep_next=True)
        _style_run(paragraph.add_run(f"{self.section_no:02d}"), size=10, bold=True, color=AMBER)
        _style_run(paragraph.add_run("   "), size=10)
        _style_run(paragraph.add_run(title.upper()), size=12, bold=True, color=INK, spacing=24)
        _paragraph_border(paragraph, "bottom", LINE_HEX, size=6, space=4)

    def summary(self):
        self.heading(self.t["sections"]["summary"])
        _para(self.doc, self.t["summary"], size=10.5, line=1.3, after=2)

    def skills(self):
        self.heading(self.t["sections"]["skills"])
        table = self.doc.add_table(rows=0, cols=2)
        _remove_table_borders(table)
        for label, value in self.t["skills"]:
            row = table.add_row()
            left, right = row.cells
            _set_cell_margins(left, 30, 30, 0, 100)
            _set_cell_margins(right, 30, 30, 100, 0)
            _style_run(_clear_cell(left).add_run(label), size=10, bold=True, color=INK)
            _style_run(_clear_cell(right).add_run(value), size=10, color=TEXT)
        _set_col_widths(table, [Cm(3.2), CONTENT_WIDTH - Cm(3.2)])

    def _period(self, exp: ExportExperience) -> str:
        months = self.t["months"]

        def fmt(value: date) -> str:
            return f"{months[value.month - 1]} {value.year}"

        end = self.t["present"] if exp.is_current or not exp.end_date else fmt(exp.end_date)
        return f"{fmt(exp.start_date)} — {end}"

    def experience(self):
        if not self.data.experiences:
            return
        self.heading(self.t["sections"]["experience"])
        table = self.doc.add_table(rows=0, cols=2)
        _remove_table_borders(table)
        for exp in self.data.experiences:
            row = table.add_row()
            _row_cant_split(row)
            left, right = row.cells
            _set_cell_margins(left, 80, 140, 0, 120)
            _set_cell_margins(right, 80, 140, 120, 0)

            period = _clear_cell(left)
            _style_run(period.add_run(self._period(exp)), size=9, bold=True, color=AMBER if exp.is_current else MUTED)

            title = _clear_cell(right)
            title.paragraph_format.space_after = Pt(1)
            _style_run(title.add_run(exp.position), size=11.5, bold=True, color=INK)
            _para(right, exp.company, size=10, color=AMBER, after=4)
            if exp.responsibilities:
                _para(right, exp.responsibilities, size=9.5, italic=True, color=MUTED, after=4)
            if exp.description:
                _para(right, exp.description, size=10, line=1.25, after=0)
        _set_col_widths(table, [Cm(3.6), CONTENT_WIDTH - Cm(3.6)])

    def projects(self):
        if not self.data.projects:
            return
        self.heading(self.t["sections"]["projects"])
        for index, project in enumerate(self.data.projects):
            title = _para(self.doc, before=10 if index else 2, after=2, keep_next=True)
            _paragraph_border(title, "left", AMBER_HEX, size=18, space=8)
            title.paragraph_format.left_indent = Cm(0.25)
            _style_run(title.add_run(project.title), size=11.5, bold=True, color=INK)

            if project.short_description:
                lead = _para(self.doc, project.short_description, size=10, italic=True, color=MUTED,
                             after=4, keep_next=True)
                lead.paragraph_format.left_indent = Cm(0.55)
            for chunk in (project.description or "").split("\n\n"):
                if chunk.strip():
                    body = _para(self.doc, chunk.strip(), size=10, line=1.25, after=4)
                    body.paragraph_format.left_indent = Cm(0.55)
            if project.stack:
                stack = _para(self.doc, after=2)
                stack.paragraph_format.left_indent = Cm(0.55)
                _style_run(stack.add_run(f"{self.t['stack']}:  "), size=9, bold=True, color=AMBER)
                _style_run(stack.add_run("  ·  ".join(s.strip() for s in project.stack.split(",") if s.strip())),
                           size=9, color=MUTED)

    def mentoring(self):
        if not (self.data.metrics or self.data.achievements):
            return
        self.heading(self.t["sections"]["mentoring"])
        if self.data.mentoring_lead:
            _para(self.doc, self.data.mentoring_lead, size=10.5, line=1.3, after=8)

        if self.data.metrics:
            count = len(self.data.metrics)
            table = self.doc.add_table(rows=1, cols=count)
            _remove_table_borders(table)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            gap_width = (CONTENT_WIDTH - Cm(0.3) * (count - 1)) / count
            for cell, metric in zip(table.rows[0].cells, self.data.metrics):
                _set_cell_shading(cell, TINT_HEX)
                _set_cell_margins(cell, 140, 140, 120, 120)
                value = _clear_cell(cell)
                value.alignment = WD_ALIGN_PARAGRAPH.CENTER
                value.paragraph_format.space_after = Pt(0)
                _style_run(value.add_run(metric.value), size=22, bold=True, color=CORAL)
                _para(cell, metric.label, size=9, color=MUTED, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
            _set_col_widths(table, [int(gap_width)] * count)
            _para(self.doc, after=4)

        if self.data.achievements:
            table = self.doc.add_table(rows=0, cols=2)
            _remove_table_borders(table)
            for item in self.data.achievements:
                row = table.add_row()
                _row_cant_split(row)
                left, right = row.cells
                _set_cell_margins(left, 50, 90, 0, 120)
                _set_cell_margins(right, 50, 90, 120, 0)
                year = _clear_cell(left)
                _style_run(year.add_run(str(item.year) if item.year else "—"), size=9.5, bold=True, color=MUTED)

                title = _clear_cell(right)
                title.paragraph_format.space_after = Pt(1)
                _style_run(title.add_run(item.title), size=10, bold=True, color=INK)
                if item.result:
                    _style_run(title.add_run(f"  —  {item.result}"), size=10, bold=True, color=AMBER)

                meta = [
                    self.t["categories"].get(item.category, item.category),
                    self.t["scopes"].get(item.scope, item.scope),
                ]
                if item.organization:
                    meta.insert(0, item.organization)
                _para(right, "  ·  ".join(meta), size=8.5, color=MUTED, after=1)
                if item.description:
                    _para(right, item.description, size=9.5, line=1.2, after=0)
            _set_col_widths(table, [Cm(1.6), CONTENT_WIDTH - Cm(1.6)])

    def build(self) -> bytes:
        self.header()
        self.summary()
        self.skills()
        self.experience()
        self.projects()
        self.mentoring()
        buffer = io.BytesIO()
        self.doc.save(buffer)
        return buffer.getvalue()


def build_portfolio_docx(data: PortfolioExportData) -> bytes:
    return _Builder(data).build()
